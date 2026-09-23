"""
download_music - coeur metier telechargement YouTube -> MP3
Utilisation module (app web):
    from download_music import load_tracks, load_tracks_from_text, download_track
Debug CLI (l'app web suffit en usage normal):
    python download_music.py
    python download_music.py fichier.xlsx --quality 320
"""

import argparse
import csv
import difflib
import logging
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

import yt_dlp

BASE_DIR = Path(__file__).parent
DEFAULT_DOWNLOADS = BASE_DIR / "downloads"
DEFAULT_LOG = BASE_DIR / "download_failures.log"

# Sous ce seuil, le MP3 est considéré comme corrompu/vide (OK silencieux + skip fantôme).
MIN_MP3_BYTES = 50_000


def is_valid_mp3(path: Path) -> bool:
    try:
        return path.is_file() and path.stat().st_size >= MIN_MP3_BYTES
    except OSError:
        return False


def sanitize_filename(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", name)
    name = re.sub(r"\s+", " ", name).strip().strip(" .")
    return name[:150] if len(name) > 150 else name


def find_spreadsheet(base: Path) -> Path | None:
    for ext in ("*.xlsx", "*.csv"):
        files = list(base.glob(ext))
        if files:
            return files[0]
    return None


def detect_columns(headers: list[str]) -> tuple[int, int]:
    ai = ti = None
    for i, h in enumerate(headers):
        low = str(h or "").lower().strip()
        if ai is None and ("artiste" in low or low == "artist"):
            ai = i
        if ti is None and ("titre" in low or low == "title" or "titre du morceau" in low):
            ti = i
    if ai is None or ti is None:
        raise ValueError(f"Colonnes Artiste/Titre introuvables. Colonnes detectees: {headers}")
    return ai, ti


def _clean(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("", "nan", "none", "null") else s


def load_tracks(path: Path) -> list[tuple[str, str]]:
    suf = path.suffix.lower()
    if suf == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
    elif suf == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    else:
        raise ValueError(f"Format non supporté ({suf}). Utilise .xlsx, .csv ou .txt.")
    if not rows:
        return []
    ai, ti = detect_columns([str(h or "") for h in rows[0]])
    out = []
    for r in rows[1:]:
        if len(r) <= max(ai, ti):
            continue
        a, t = _clean(r[ai]), _clean(r[ti])
        if a and t:
            out.append((a, t))
    return out


def load_tracks_from_text(text: str) -> list[tuple[str, str]]:
    tracks = []
    for line in text.strip().splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        sep = None
        for cand in [" - ", " – ", " — ", " | ", ";"]:
            if cand in line:
                sep = cand
                break
        if sep:
            a, t = line.split(sep, 1)
        elif "," in line and line.count(",") == 1:
            a, t = line.split(",", 1)
        else:
            continue
        a, t = a.strip(), t.strip()
        a = a.lstrip("-–— ").strip()  # ex. "-Prey, Emrld!" -> "Prey, Emrld!"
        # Restes de noms de fichiers : "Testosterone.Exe" -> "Testosterone"
        t = re.sub(r"\.(mp3|mp4|m4a|wav|flac|ogg|opus|webm|mov|avi|mkv|exe)$", "", t, flags=re.IGNORECASE).strip()
        if a and t:
            tracks.append((a, t))
    return tracks


def _norm(s: str) -> str:
    s = s.lower()
    s = re.sub(r"[^\w\s]", " ", s, flags=re.UNICODE)
    return re.sub(r"\s+", " ", s).strip()


# Mots signalant un format long légitime (podcasts, mixes, lives voulus...)
_LONG_WORDS = ("mix", "set", "compilation", "hour", "heures", "live", "podcast",
               "episode", "épisode", "émission", "documentaire", "interview")


def clean_error(e) -> str:
    """Erreurs yt-dlp lisibles : strip codes couleur ANSI (+ résidus sans ESC)."""
    msg = re.sub(r"\x1b\[[0-9;]*[a-zA-Z]", "", str(e))
    msg = re.sub(r"\[[0-9;]*m", "", msg)
    return re.sub(r"\s+", " ", msg).strip()


def friendly_error(e, limit: int = 300) -> str:
    """Erreur courte et actionnable pour l'UI (le message brut yt-dlp fait 10 lignes)."""
    m = clean_error(e)
    if "not a bot" in m or "Sign in to confirm" in m:
        return ("YouTube bloque (anti-bot). Ajoute un cookies.txt YouTube à côté de app.py "
                "(voir README) ou relance avec --cookies-from-browser brave, puis relance la liste.")
    return m if len(m) <= limit else m[:limit] + "…"
# Pénalisés seulement si absents de la demande (protège les "(Slowed)", lives voulus, etc.)
_SUSPECT = ("remix", "bootleg", "cover", "reprise", "live", "concert", "karaok",
            "instrumental", "sped up", "mashup", "parodie", "parody", "reaction",
            "tutoriel", "tutorial", "interview", "8d ", "8d audio")


def pick_best_video(entries: list[dict], artiste: str, titre: str) -> dict | None:
    """Choisit le meilleur candidat parmi les résultats de recherche (stdlib uniquement)."""
    query = _norm(f"{artiste} {titre}")
    qtokens = query.split()
    best, best_score = None, -1.0
    for e in entries:
        if not isinstance(e, dict):
            continue
        vtitle = _norm(str(e.get("title") or ""))
        if not vtitle:
            continue
        vtokens = vtitle.split()
        vset = set(vtokens)
        matched = sum(1 for t in qtokens if t in vset)
        recall = matched / max(len(qtokens), 1)
        precision = matched / max(len(vtokens), 1)
        f1 = 2 * recall * precision / (recall + precision) if (recall + precision) else 0.0
        sim = difflib.SequenceMatcher(None, query, vtitle).ratio()
        score = 2.0 * f1 + 1.0 * sim
        if _norm(artiste) and _norm(artiste) in vtitle:
            score += 0.5
        # On télécharge de l'audio : bonus piste audio officielle, malus clips
        # (souvent intros/dialogues). Aligné avec notre requête "Official Audio".
        if "official audio" in vtitle:
            score += 0.6
        if "music video" in vtitle or "official video" in vtitle or "clip officiel" in vtitle:
            score -= 0.4
        for bad in _SUSPECT:
            if bad in vtitle and bad not in query:
                score -= 1.0
        dur = e.get("duration") or 0
        long_ok = any(w in query for w in _LONG_WORDS)
        if dur and dur > 1800 and not long_ok:
            continue  # compilations / lives de 30 min+ : jamais le morceau
        if dur and (dur < 30 or dur > 900) and not long_ok:
            score *= 0.5
        if score > best_score:
            best, best_score = e, score
    return best


_AUTH = {}  # ex. {"cookiefile": "..."} ou {"cookiesfrombrowser": ("brave",)}


def set_auth(cookiefile: str | None = None, from_browser: str | None = None) -> None:
    """Auth YouTube anti-bot. Priorité : --cookies > --cookies-from-browser > cookies.txt auto."""
    global _AUTH
    _AUTH = {}
    if cookiefile:
        _AUTH["cookiefile"] = str(cookiefile)
    elif from_browser:
        _AUTH["cookiesfrombrowser"] = (from_browser,)


def _cookie_opts() -> dict:
    cfile = BASE_DIR / "cookies.txt"
    try:
        if cfile.is_file() and cfile.stat().st_size > 0:
            return {"cookiefile": str(cfile)}
    except OSError:
        pass
    return {}


def _base_opts(extra: dict | None = None) -> dict:
    # retries + pauses : YouTube renvoie 503 / anti-bot quand on enchaîne les requêtes
    opts = {
        "quiet": True,
        "no_warnings": True,
        "noplaylist": True,
        "retries": 10,
        "sleep_requests": 1,
        "extractor_args": {"youtube": {"player_client": ["android", "web"]}},
    }
    opts.update(_cookie_opts())
    opts.update(_AUTH)
    if extra:
        opts.update(extra)
    return opts


def search_candidates(query: str, n: int = 8) -> list[dict]:
    # Recherche légère : peu d'essais (le téléchargement, lui, insiste avec retries=10).
    # Une recherche reste un appel bloquant : l'annulation est honorée entre les requêtes.
    opts = _base_opts({"skip_download": True, "retries": 3})
    with yt_dlp.YoutubeDL(opts) as ydl:
        info = ydl.extract_info(f"ytsearch{n}:{query}", download=False)
    entries = info.get("entries", []) if isinstance(info, dict) else []
    return [e for e in (list(entries) if entries else []) if isinstance(e, dict)]


def _first_sane(entries: list[dict], artiste: str, titre: str) -> dict | None:
    """Dernier recours : 1er résultat titré, durée raisonnable (jamais une compil 30 min+)."""
    query = _norm(f"{artiste} {titre}")
    long_ok = any(w in query for w in _LONG_WORDS)
    for e in entries:
        if not isinstance(e, dict) or not _norm(str(e.get("title") or "")):
            continue
        dur = e.get("duration") or 0
        if dur and dur > 1800 and not long_ok:
            continue
        return e
    return None


def _find_best(artiste: str, titre: str, abort: "callable[[], bool] | None" = None) -> tuple[dict | None, str, bool]:
    """Essaie des requêtes de moins en moins strictes. Retourne (video, requête, résultats_vus)."""
    queries = list(dict.fromkeys(
        q for q in (f"{artiste} - {titre} Official Audio", f"{artiste} - {titre}", titre) if q.strip()
    ))
    seen_any = False
    for q in queries:
        if abort is not None and abort():
            raise Aborted()
        try:
            entries = search_candidates(q)
        except Exception as e:
            logging.warning(f"[SEARCH] {q} | {clean_error(e)}")
            continue
        if not entries:
            continue
        seen_any = True
        best = pick_best_video(entries, artiste, titre) or _first_sane(entries, artiste, titre)
        if best:
            return best, q, True
    return None, queries[0], seen_any


def download_track(artiste: str, titre: str, downloads: Path, quality: str,
                   abort: "callable[[], bool] | None" = None) -> bool:
    """Telecharge un morceau. Si abort() devient vrai, leve Aborted (partiels nettoyés)."""
    filename = sanitize_filename(f"{artiste} - {titre}")
    expected_mp3 = downloads / f"{filename}.mp3"
    if is_valid_mp3(expected_mp3):
        logging.info(f"[SKIP] {filename} deja present")
        return True
    if expected_mp3.exists():
        expected_mp3.unlink()  # fichier corrompu/vide : on re-télécharge au lieu de skipper
    best, used_query, seen_any = _find_best(artiste, titre, abort=abort)
    if abort is not None and abort():
        raise Aborted()
    if not best:
        if seen_any:
            raise RuntimeError(f"aucun candidat valable pour : {artiste} - {titre}")
        raise RuntimeError(f"aucun resultat YouTube pour : {artiste} - {titre}")
    vid = best.get("id") or ""
    video_url = best.get("webpage_url") or (f"https://www.youtube.com/watch?v={vid}" if vid else "")
    if not video_url:
        raise RuntimeError(f"URL introuvable pour : {artiste} - {titre}")
    logging.info(f"[MATCH] {filename} -> {best.get('title')} (via « {used_query} »)")
    ydl_opts = _base_opts({
        "format": "bestaudio/best",
        "outtmpl": str(downloads / f"{filename.replace('%', '%%')}.%(ext)s"),
        "postprocessors": [{"key": "FFmpegExtractAudio", "preferredcodec": "mp3", "preferredquality": quality}],
        "sleep_interval": 1,
        "max_sleep_interval": 3,
    })
    if abort is not None:
        def _hook(d):
            if abort():
                raise Aborted()
        ydl_opts["progress_hooks"] = [_hook]
    try:
        with yt_dlp.YoutubeDL(ydl_opts) as ydl:
            ydl.download([video_url])
    except Aborted:
        _clean_partials(downloads, filename)
        raise
    except Exception as e:
        raise RuntimeError(clean_error(e))
    if not is_valid_mp3(expected_mp3):
        if expected_mp3.exists():
            expected_mp3.unlink()
        raise RuntimeError(f"fichier MP3 invalide ou vide apres telechargement : {filename}")
    return True


class Aborted(Exception):
    """Levé quand abort() interrompt un téléchargement en cours (bouton Arrêter)."""


def _clean_partials(downloads: Path, filename: str) -> None:
    """Supprime les restes d'un téléchargement interrompu (jamais un MP3 valide)."""
    for f in downloads.iterdir():
        if not f.is_file() or not f.name.startswith(filename + "."):
            continue
        try:
            if f.suffix == ".mp3" and is_valid_mp3(f):
                continue
            f.unlink()
        except OSError:
            pass


def main():
    parser = argparse.ArgumentParser(description="Telecharge les audios YouTube depuis un tableur ou une liste")
    parser.add_argument("fichier", nargs="?", help="Chemin vers .xlsx/.csv (auto-detecte sinon)")
    parser.add_argument("--output", "-o", default=str(DEFAULT_DOWNLOADS), help="Dossier de sortie")
    parser.add_argument("--quality", default="192", help="Bitrate MP3 (defaut: 192)")
    parser.add_argument("--log", default=str(DEFAULT_LOG), help="Fichier de log des echecs")
    parser.add_argument("--cookies", default=None, help="Fichier cookies.txt YouTube (anti-bot)")
    parser.add_argument("--cookies-from-browser", default=None, help="Ex. brave, chrome, firefox (anti-bot)")
    args = parser.parse_args()
    set_auth(args.cookies, args.cookies_from_browser)
    if args.quality not in ("192", "256", "320"):
        print(f"Qualité invalide ({args.quality}), 192/256/320 acceptés — 192 utilisée.")
        args.quality = "192"
    if args.fichier:
        source = Path(args.fichier)
    else:
        source = find_spreadsheet(BASE_DIR)
        if not source:
            print("Aucun .xlsx/.csv trouve dans", BASE_DIR)
            print("Usage: python download_music.py chemin/vers/fichier.xlsx")
            sys.exit(1)
    if not source.exists():
        print(f"Fichier introuvable: {source}")
        sys.exit(1)
    downloads = Path(args.output)
    downloads.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        handlers=[logging.FileHandler(args.log, encoding="utf-8"), logging.StreamHandler()],
    )
    try:
        tracks = load_tracks(source)
    except Exception as e:
        logging.error(f"Erreur lecture tableur: {e}")
        sys.exit(1)
    logging.info(f"{len(tracks)} morceaux trouves dans {source.name}")
    logging.info(f"Sortie: {downloads} | Qualite: {args.quality} kbps")
    success = skipped = failed = 0
    for artiste, titre in tracks:
        filename = sanitize_filename(f"{artiste} - {titre}")
        expected = downloads / f"{filename}.mp3"
        if is_valid_mp3(expected):
            logging.info(f"[SKIP] {filename}")
            skipped += 1
            continue
        try:
            download_track(artiste, titre, downloads, args.quality)
            logging.info(f"[OK] {filename}")
            success += 1
        except Exception as e:
            err = friendly_error(e)
            logging.error(f"[FAIL] {filename} | {err}")
            failed += 1
    logging.info(f"Termine: {success} telecharges, {skipped} deja presents, {failed} echecs sur {len(tracks)}")


if __name__ == "__main__":
    main()
